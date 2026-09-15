import ast
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app/src/main/python"))
from mobile_api import inspect_file, process_batch, validate_options


def reference_processor():
    """Load the original processing code without importing the desktop GUI."""
    tree = ast.parse((ROOT / "reference/excel_processor.py").read_text(encoding="utf-8-sig"))
    nodes = []
    for node in tree.body:
        if isinstance(node, ast.Import):
            if all(not alias.name.startswith("tkinter") for alias in node.names):
                nodes.append(node)
        elif isinstance(node, ast.ImportFrom) and not (node.module or "").startswith("tkinter"):
            nodes.append(node)
        elif isinstance(node, (ast.FunctionDef, ast.ClassDef)) and node.name in (
                "question_sort_key", "detect_id_segments", "ExcelProcessor"):
            nodes.append(node)
    namespace = {}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), "desktop_reference", "exec"), namespace)
    return namespace["ExcelProcessor"]


def fixture(path, offset=0, duplicate_header=False):
    columns = ["PINFL", "FIO", "Savol ID", "Savol Tr", "Yuklangan variant", "Natija"]
    rows = []
    # Synthetic identifiers only: two complete participants, one incomplete.
    for person, questions in [(30000000000001, [1, 2, 10]), (40000000000002, [1, 2, 10]), (30000000000003, [1])]:
        for question in questions:
            rows.append([str(person), "Тест " + str(person)[-1], question + offset,
                         question, {1: "a", 2: "B", 10: "D"}[question], 0 if question == 2 else 1])
    if duplicate_header:
        rows.insert(0, columns)
    df = pd.DataFrame(rows, columns=columns)
    if path.suffix == ".csv": df.to_csv(path, index=False)
    else: df.to_excel(path, index=False)


def workbook_contents(path):
    book = load_workbook(io.BytesIO(Path(path).read_bytes()))
    contents = []
    for sheet in book:
        contents.append((sheet.title, [
            [(c.value, c.font.bold, c.border.left.style, c.number_format) for c in row]
            for row in sheet.iter_rows()
        ], {k: v.width for k, v in sheet.column_dimensions.items()}))
    book.close()
    return contents


class ProcessingTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self): self.tmp.cleanup()

    def batch(self, files, **settings):
        options = dict(sort_column="Savol ID", header_column="Savol Tr", **settings)
        return json.loads(process_batch(json.dumps(files), json.dumps(options), str(self.root / "results")))

    def test_matches_desktop_cells_styles_scores_errors_and_shifted_groups(self):
        Reference = reference_processor()
        for offset, extension in [(0, ".xlsx"), (200, ".csv")]:
            with self.subTest(offset=offset):
                source = self.root / ("input" + extension)
                fixture(source, offset, duplicate_header=True)
                options = dict(add_gender=True, add_scores=True, multiplier=2.5,
                               size_rules=[{"count": 2, "prefix": "K-"}, {"count": 1, "prefix": "PM-"}])
                report = self.batch([{"name": source.name, "path": str(source)}], **options)
                expected = self.root / "expected.xlsx"
                Reference(log_callback=lambda _: None, sort_column="Savol ID", header_column="Savol Tr", **options).process(str(source), str(expected))
                row = report["results"][0]
                self.assertEqual(report["success_count"], 1)
                self.assertEqual(row["participants"], 2)
                self.assertEqual(row["error_rows"], 1)
                self.assertEqual(workbook_contents(row["path"]), workbook_contents(expected))
                book = load_workbook(io.BytesIO(Path(row["path"]).read_bytes()))
                self.assertEqual(book.sheetnames, ["Данные", "Ошибки"])
                self.assertEqual([book.active.cell(2, c).value for c in (4, 5, 6)], ["A", "B", "D"])
                self.assertEqual([book.active.cell(2, c).value for c in (10, 11)], [5, 2])
                book.close()

    def test_range_prefixes_separate_key_match_reference(self):
        source = self.root / "input.xlsx"; fixture(source, 100)
        options = dict(prefix_column="Savol Tr", suffixes=[{"prefix": "PM-", "from": 1, "to": 2}])
        report = self.batch([{"name": source.name, "path": str(source)}], **options)
        expected = self.root / "expected.xlsx"
        reference_processor()(log_callback=lambda _: None, sort_column="Savol ID", header_column="Savol Tr", **options).process(str(source), str(expected))
        self.assertEqual(workbook_contents(report["results"][0]["path"]), workbook_contents(expected))

    def test_archive_contains_every_successful_file_with_unique_names(self):
        source = self.root / "input.csv"; fixture(source)
        files = [{"name": name, "path": str(source)} for name in ("тест.xlsx", "тест.csv", "ТЕСТ.xlsx")]
        files.append({"name": "broken.xlsx", "path": str(self.root / "missing.xlsx")})
        report = self.batch(files)
        self.assertEqual(report["success_count"], 3)
        self.assertTrue(report["results"][-1]["error"])
        with zipfile.ZipFile(report["archive"]) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), 3)
            self.assertEqual(len(set(n.casefold() for n in names)), 3)
            self.assertIsNone(archive.testzip())
            for row in report["results"][:3]:
                self.assertEqual(archive.read(Path(row["path"]).name), Path(row["path"]).read_bytes())
        saved = json.loads((self.root / "results/latest.json").read_text(encoding="utf-8"))
        self.assertEqual(saved, report)

    def test_empty_and_failed_files_are_never_reported_as_successful(self):
        source = self.root / "empty.csv"
        source.write_text("PINFL,FIO,Savol ID,Yuklangan variant,Natija\n", encoding="utf-8")
        report = self.batch([{"name": source.name, "path": str(source)}])
        self.assertEqual(report["success_count"], 0)
        self.assertEqual(report["archive"], "")
        self.assertTrue(report["results"][0]["error"])

    def test_repeated_runs_preserve_previous_outputs(self):
        source = self.root / "input.csv"; fixture(source)
        items = [{"name": source.name, "path": str(source)}]
        first = self.batch(items); original = Path(first["archive"]).read_bytes()
        second = self.batch(items)
        self.assertNotEqual(first["archive"], second["archive"])
        self.assertEqual(Path(first["archive"]).read_bytes(), original)

    def test_inspect_and_invalid_settings(self):
        source = self.root / "input.csv"; fixture(source)
        info = json.loads(inspect_file(str(source), "Savol ID"))
        self.assertEqual([s["count"] for s in info["segments"]], [2, 1])
        self.assertIn("PINFL", info["headers"])
        for bad in (-1, float("nan"), float("inf")):
            with self.assertRaises(ValueError): validate_options({"multiplier": bad})


if __name__ == "__main__": unittest.main()
